"""
从参考报价表生成「技术规格书」。

流程：
    参考报价表里的型号（如 GKV225）
        → ① 从**产品资料**里检索该型号的技术参数原文
        → ② 让本地模型把原文抽成结构化参数（JSON）
        → ③ 汇总成 Excel 规格书

⚠️ 铁律：**只抽原文里写了的，抽不到就写「资料未载明」**。
   阀门参数写错是要出事故的，宁可留白，绝不推测。

用法：
    python src/quote/make_spec.py --prices "data/参考报价_Q260509-10_Vanko.xlsx"
    python src/quote/make_spec.py --series GKV225,GKV-243,GKV326
    python src/quote/make_spec.py --prices "..." --limit 5      # 先试 5 个系列
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from pathlib import Path

import chromadb
import ollama
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.config import CHROMA_DIR, EMBED_MODEL, LLM_MODEL  # noqa: E402

PROJECT = Path(__file__).resolve().parents[2]
CACHE = PROJECT / "data" / "spec_cache"
# 检索策略 / 提示词一旦升级，就把版本号 +1 —— 旧缓存自动作废，不必手动删。
CACHE_VER = "v5"

# 规格书的参数表头（也是给模型的键名）
FIELDS = [
    "阀门类型", "结构形式", "口径范围", "压力等级", "温度范围",
    "阀体材质", "设计标准", "试验标准", "连接方式", "操作方式", "备注",
]

PROMPT = """你是阀门技术资料整理助手。下面是从 该企业 产品资料里检索到的原文片段（可能含 OCR 噪声）。

请**严格依据原文**，抽取型号 **{series}** 的技术参数，输出 JSON。

规则：
1. ⚠️ **只抽 {series} 这一个型号**。同一本手册里往往并列着很多型号，
   那些不属于 {series} 的内容**一律不要**；无法确定某段参数是否属于 {series} 的，填 "资料未载明"。
2. 只填原文中**明确写到**的内容。原文没写的，值一律填 "资料未载明"，**绝对不要推测、不要编造**。
3. 参数值尽量保留原文写法（如 "PN16/40"、"-20 to 180"、"CF8M/CF8/WCB"）。
4. OCR 噪声可做常识性修正（如 "TEMPERATURERANGE:-20to 180C" → "-20 to 180 ℃"），
   但不得增加原文里没有的信息。
5. 必须输出**严格的 JSON**，键固定为下列这些，一个不多一个不少。

键：{keys}

--- 原文开始 ---
{context}
--- 原文结束 ---
"""


# ---------------------------------------------------------------- 检索

def _variant_re(series: str):
    """型号可能有多种写法：GKV326 / GKV-326 / GKV 326。生成一个宽松匹配的正则。"""
    m = re.match(r"([A-Za-z]+)[\s\-]*(\d+)([A-Za-z0-9]*)", series)
    if not m:
        return None
    head, num, tail = m.group(1), m.group(2), m.group(3)
    return re.compile(re.escape(head) + r"[\s\-]*" + re.escape(num) + re.escape(tail), re.I)


# 参数关键词：用来从"同文件其他页"里挑出真正载有参数的切片
PARAM_HINT = re.compile(
    r"PN\s*\d|\d+\s*#|\d+\s*LB|CLASS|MPa|psi|温度|temperature|材质|material|"
    r"API|ASME|EN\s*\d|GB/T|DN\s*\d|设计标准|试验标准|连接|操作方式", re.I)


def _variants(series: str) -> list[str]:
    """型号在资料里的几种写法：GKV326 / GKV-326 / GKV 326。"""
    m = re.match(r"([A-Za-z]+)[\s\-]*(\d+)([A-Za-z0-9]*)", series)
    if not m:
        return [series]
    h, n, t = m.group(1), m.group(2), m.group(3)
    out = [f"{h}{n}{t}", f"{h}-{n}{t}", f"{h} {n}{t}"]
    return list(dict.fromkeys(out))


def retrieve(series: str, k: int = 6) -> list[tuple[str, str]]:
    """
    捞取该型号的参数原文。

    ⚠️ 顺序很重要：**先用"精确包含型号"捞，再用语义检索补**。
    因为型号是**精确标识**，而语义检索对它并不敏感 ——
    实测踩过：查 GKV326（闸阀）却召回 GKV225（球阀）的段落，
    结果把闸阀写成了球阀。这类错误在技术规格书里是致命的。
    """
    col = chromadb.PersistentClient(path=str(CHROMA_DIR)).get_collection("knowledge")
    base_where = {"$and": [{"status": {"$eq": "active"}},
                           {"top": {"$eq": "产品资料"}}]}
    cands: list[tuple[str, str]] = []
    seen: set[tuple[str, str]] = set()

    # ① 精确包含型号（最可靠）
    for v in _variants(series):
        try:
            d = col.get(where=base_where,
                        where_document={"$contains": v},
                        limit=12, include=["documents", "metadatas"])
        except Exception:
            continue
        for doc, m in zip(d.get("documents") or [], d.get("metadatas") or []):
            key = (m.get("source", ""), doc[:60])
            if key in seen:
                continue
            seen.add(key)
            cands.append((doc, m.get("source", "")))

    # ② 不够时用语义检索补（用于资料里写法差异较大的情况）
    if len(cands) < k:
        queries = [
            f"{series} valve specification design standard pressure temperature material",
            f"{series} 阀门 设计标准 压力 温度 材质 连接",
        ]
        for q in queries:
            qv = ollama.embed(model=EMBED_MODEL, input=[q])["embeddings"][0]
            res = col.query(query_embeddings=[qv], n_results=20, where=base_where)
            rx = _variant_re(series)
            for doc, m in zip(res["documents"][0], res["metadatas"][0]):
                if rx and not rx.search(doc):     # 语义召回也要过型号这一关
                    continue
                key = (m.get("source", ""), doc[:60])
                if key in seen:
                    continue
                seen.add(key)
                cands.append((doc, m.get("source", "")))
            if len(cands) >= k:
                break

    # ③ 同文件补充：参数常分散在同一本手册的不同页
    #    （命中的可能只是目录页或封面，真正的参数表在另一页）
    for s in list(dict.fromkeys(x[1] for x in cands))[:3]:
        if len(cands) >= k + 4:
            break
        try:
            d = col.get(where={"source": {"$eq": s}}, limit=40,
                        include=["documents", "metadatas"])
        except Exception:
            continue
        for doc, m in zip(d.get("documents") or [], d.get("metadatas") or []):
            if not PARAM_HINT.search(doc):        # 只保留确实载有参数的页
                continue
            key = (m.get("source", ""), doc[:60])
            if key in seen:
                continue
            seen.add(key)
            cands.append((doc, m.get("source", "")))

    return cands[: k + 4]


# ---------------------------------------------------------------- 抽取

def extract_params(series: str, chunks: list[tuple[str, str]]) -> dict:
    """调本地模型把原文抽成结构化参数（带缓存）。"""
    CACHE.mkdir(parents=True, exist_ok=True)
    cf = CACHE / f"{re.sub(r'[^A-Za-z0-9]', '_', series)}_{CACHE_VER}.json"
    if cf.exists():
        try:
            return json.loads(cf.read_text(encoding="utf-8"))
        except Exception:
            pass

    ctx = "\n\n".join(f"[{i}] 来源：{s}\n{d[:1500]}"
                      for i, (d, s) in enumerate(chunks, 1))
    try:
        r = ollama.chat(
            model=LLM_MODEL,
            messages=[{"role": "user",
                       "content": PROMPT.format(series=series,
                                               keys="、".join(FIELDS), context=ctx)}],
            format="json",
            options={"temperature": 0.1},
        )
        raw = r["message"]["content"]
        data = json.loads(raw)
    except Exception as e:
        data = {k: "抽取失败" for k in FIELDS}
        data["备注"] = f"{type(e).__name__}: {str(e)[:60]}"

    out = {k: str(data.get(k, "资料未载明")).strip() or "资料未载明" for k in FIELDS}
    out["_sources"] = sorted({s for _, s in chunks})
    cf.write_text(json.dumps(out, ensure_ascii=False, indent=1), encoding="utf-8")
    return out


# ---------------------------------------------------------------- 出表

HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(color="FFFFFF", bold=True, size=10)
CELL = Font(size=10)


def write_spec(rows: list[dict], out_path: Path, src_note: str) -> None:
    wb = Workbook()
    ws = wb.active
    ws.title = "技术规格书"

    heads = ["我方型号"] + FIELDS + ["参数依据（产品资料）"]
    widths = [14] + [22] * len(FIELDS) + [46]
    ws.append(heads)
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        c = ws.cell(row=1, column=i)
        c.fill, c.font = HDR_FILL, HDR_FONT
        c.alignment = Alignment(vertical="center", horizontal="center", wrap_text=True)
    ws.row_dimensions[1].height = 26
    ws.freeze_panes = "B2"

    for r in rows:
        ws.append([r["series"]] + [r.get(k, "资料未载明") for k in FIELDS]
                  + ["; ".join(r.get("_sources", []))[:80]])
        i = ws.max_row
        for j in range(1, len(heads) + 1):
            c = ws.cell(row=i, column=j)
            c.font = CELL
            c.alignment = Alignment(vertical="top", wrap_text=True)
        ws.row_dimensions[i].height = 34

    # 说明页
    ws2 = wb.create_sheet("说明")
    notes = [
        ["本规格书由程序从产品资料中自动抽取，供内部参考与对外初稿，"], 
        ["**对外发出前必须由技术部门复核**。"],
        [""],
        ["抽取规则"],
        ["· 只填产品资料中明确写到的内容；没写的一律标「资料未载明」，不做任何推测。"],
        ["· OCR 噪声做常识性修正，但不增加原文没有的信息。"],
        ["· 「参数依据」列列出该型号参数取自哪些产品资料，便于核对。"],
        [""],
        ["数据来源"],
        [f"型号清单：{src_note}"],
        ["参数来源：知识库中「产品资料」目录（D:\\RAG资料\\产品资料）"],
        ["模型：本地 " + LLM_MODEL + "（数据不出本机）"],
    ]
    for i, row in enumerate(notes, 1):
        c = ws2.cell(row=i, column=1, value=row[0])
        c.font = Font(bold=True, size=11) if row[0] in (
            "抽取规则", "数据来源") or row[0].startswith("**") else CELL
        c.alignment = Alignment(wrap_text=False)
    ws2.column_dimensions["A"].width = 92

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)


# ---------------------------------------------------------------- 入口

def series_from_prices(path: Path) -> list[str]:
    """从参考报价表里取「系列」列，去重保序。"""
    wb = load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if not rows:
        return []
    head = [str(c or "").strip() for c in rows[0]]
    j = 2
    for i, h in enumerate(head):
        if "系列" in h or "参考型号" in h:
            j = i
            break
    out: list[str] = []
    for r in rows[1:]:
        if j >= len(r):
            continue
        v = str(r[j] or "").strip()
        if v and v not in ("—", "None") and v not in out:
            out.append(v)
    return out


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("--prices", default=None, help="参考报价表 xlsx（从中取型号）")
    ap.add_argument("--series", default=None, help="直接给型号系列，逗号分隔")
    ap.add_argument("--limit", type=int, default=0, help="只处理前 N 个系列（试跑用）")
    ap.add_argument("--out", default=None)
    args = ap.parse_args()

    if args.series:
        series_list = [s.strip() for s in args.series.split(",") if s.strip()]
        src_note = "命令行指定"
    elif args.prices:
        p = Path(args.prices)
        if not p.is_absolute():
            p = PROJECT / p
        series_list = series_from_prices(p)
        src_note = p.name
    else:
        print("请用 --prices 或 --series 指定型号来源")
        return

    if not series_list:
        print("没取到任何型号系列")
        return
    if args.limit:
        series_list = series_list[: args.limit]

    print(f"待生成规格书的型号系列：{len(series_list)} 个")
    print("  " + "、".join(series_list[:12]) + ("…" if len(series_list) > 12 else ""))
    print()

    rows: list[dict] = []
    for i, s in enumerate(series_list, 1):
        chunks = retrieve(s)
        if not chunks:
            print(f"[{i}/{len(series_list)}] {s}  ← 产品资料里没检索到")
            rows.append({"series": s, **{k: "资料未载明" for k in FIELDS}, "_sources": []})
            continue
        data = extract_params(s, chunks)
        rows.append({"series": s, **data})
        got = sum(1 for k in FIELDS if data.get(k) not in ("资料未载明", "抽取失败"))
        print(f"[{i}/{len(series_list)}] {s}  ← 命中 {len(chunks)} 段原文，抽到 {got}/{len(FIELDS)} 项参数")

    out = Path(args.out) if args.out else (PROJECT / "data" / "技术规格书.xlsx")
    write_spec(rows, out, src_note)
    print(f"\n技术规格书已生成：{out}")

    print("\n预览：")
    for r in rows[:6]:
        print(f"   {r['series']:<12} 类型={r.get('阀门类型','')[:12]:<14} "
              f"口径={r.get('口径范围','')[:18]:<20} 压力={r.get('压力等级','')[:14]}")


if __name__ == "__main__":
    main()
