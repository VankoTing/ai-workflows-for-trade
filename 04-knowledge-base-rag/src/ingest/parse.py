"""
解析模块：把各种格式的文件变成纯文本。

支持：PDF（pypdf）、Word（python-docx）、Excel（openpyxl）、纯文本。

直接运行可做「解析预览」：
    python src/ingest/parse.py --dir "D:/RAG资料/产品资料" --limit 12
"""
from __future__ import annotations

import argparse
import logging
import re
import sys
from pathlib import Path

# 允许直接运行时 import 到 src.config
sys.path.insert(0, str(Path(__file__).resolve().parents[2]))

from src.config import DATA_ROOT  # noqa: E402

# pypdf 遇到结构有瑕疵的 PDF 会刷一堆 "Ignoring wrong pointing object" 警告，
# 这些不影响解析结果，但会淹没输出，这里把它压到 ERROR 级。
logging.getLogger("pypdf").setLevel(logging.ERROR)


# ============================================================
# 判断"抽出来的到底是不是有用的内容"
# ============================================================
# 这些是**规格参数**专属词（带词边界），而不是 valve/ball 这类"到处都有"的词。
# 教训来自 docs/lessons.md L03：关键词选得太宽，等于没选。
SPEC_PATTERNS = [
    r"\bpressure\b", r"\btemperature\b", r"\bclass\b", r"\bpn\b", r"\bdn\b",
    r"\basme\b", r"\bapi\b", r"\bastm\b", r"\brating\b", r"\bcwp\b",
    r"\bwcb\b", r"\bcf8m\b", r"\bcf3m\b", r"\blf2\b", r"\bseat\b",
    r"\btorque\b", r"\bflange\b", r"\btest\b",
]

# 页眉页脚的特征：网址、邮箱、电话、传真
BOILERPLATE_PATTERNS = [
    r"^https?://", r"^www\.", r"@\w+\.(com|de|cn|net)", r"\btel\b", r"\bfax\b", r"e-?mail",
]


def effective_stats(text: str) -> tuple[int, int]:
    """
    评估"有效文本"的两个指标（对应 lessons.md 的 L03）：

      ① 去重后长度：页眉页脚每页重复，去重后会大幅缩水；
      ② 规格关键词命中数：这才是"有没有我要的内容"的直接证据。

    只看字符总数会被噪音骗（9708 字全是大厦地址和电话）。
    """
    lines = [l.strip() for l in text.split("\n") if l.strip()]
    uniq_len = sum(len(l) for l in set(lines))
    low = text.lower()
    hits = sum(1 for p in SPEC_PATTERNS if re.search(p, low))
    return uniq_len, hits


def strip_boilerplate(text: str, min_repeat: int = 3, max_len: int = 100) -> str:
    """
    去掉每页重复出现的页眉页脚。

    原理：一行的内容若在全文出现 ≥ min_repeat 次，且不长，那它多半是
    "每页都印一遍"的装饰性文字（公司名/地址/电话/网址），不是知识。
    这一步能显著减少噪音切片，避免它们淹没真正的知识（见 L10）。

    ⚠️ 副作用：如果某句话本来就该重复出现（如表格表头），也会被误删，
    所以阈值不要设太小。
    """
    from collections import Counter

    lines = [l.strip() for l in text.split("\n") if l.strip()]
    cnt = Counter(lines)
    out: list[str] = []
    for l in lines:
        if cnt[l] >= min_repeat and len(l) <= max_len:
            continue
        if any(re.search(p, l, re.I) for p in BOILERPLATE_PATTERNS):
            continue
        out.append(l)
    return "\n".join(out)


# ============================================================
# 各格式解析
# ============================================================
def parse_pdf(path: Path) -> dict:
    """
    解析 PDF。

    PDF 分两种，这是新手最容易踩的坑：
      ① 文本型 PDF —— 文字是可选中、可复制的，能直接抽出来；
      ② 扫描件 PDF —— 本质是"图片"，里面没有文字层，必须走 OCR。

    ⚠️ 注意：**不能只看"有没有抽出字符"**。很多产品目录的 PDF，
    文本层里只有每页重复的页眉页脚（公司名/地址/电话），真正的参数是图片。
    所以这里用 effective_stats 的两个指标来判断（见 docs/lessons.md L02/L03）。
    """
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    pages = len(reader.pages)
    texts = []
    for p in reader.pages:
        try:
            texts.append(p.extract_text() or "")
        except Exception:
            texts.append("")
    text = "\n".join(texts).strip()
    uniq_len, hits = effective_stats(text)
    # 判定逻辑（详见 docs/lessons.md 的 L02 / L03）：
    #   抽到的文本里既没有领域关键词、去重后内容又很少 → 抽的是页眉页脚噪音，必须 OCR。
    #   不再用"每页字符数"这种数量指标——它会被重复噪音骗。
    needs_ocr = (hits == 0 and uniq_len < 2000) or uniq_len < 150
    return {
        "text": text,
        "pages": pages,
        "kind": "pdf",
        "uniq_len": uniq_len,
        "hits": hits,
        "needs_ocr": needs_ocr,
    }


def parse_docx(path: Path) -> dict:
    """解析 Word（.docx）。表格内容也会被拼进来。"""
    import docx

    d = docx.Document(str(path))
    parts = [p.text for p in d.paragraphs if p.text.strip()]
    for t in d.tables:
        for row in t.rows:
            cells = [c.text.strip() for c in row.cells]
            if any(cells):
                parts.append(" | ".join(cells))
    text = "\n".join(parts).strip()
    return {"text": text, "pages": 0, "kind": "docx", "needs_ocr": False}


def parse_xlsx(path: Path) -> dict:
    """
    解析 Excel（.xlsx）。每个 sheet 转成"制表符分隔"的文本。
    注意：老的 .xls 格式 openpyxl 读不了，需要先另存为 .xlsx。
    """
    from openpyxl import load_workbook

    wb = load_workbook(str(path), read_only=True, data_only=True)
    parts = []
    for ws in wb.worksheets:
        parts.append(f"### Sheet: {ws.title}")
        for row in ws.iter_rows(values_only=True):
            cells = ["" if c is None else str(c).strip() for c in row]
            if any(cells):
                parts.append("\t".join(cells))
    text = "\n".join(parts).strip()
    return {"text": text, "pages": 0, "kind": "xlsx", "needs_ocr": False}


def parse_text(path: Path) -> dict:
    for enc in ("utf-8", "gbk", "latin-1"):
        try:
            text = path.read_text(encoding=enc)
            return {"text": text.strip(), "pages": 0, "kind": "text", "needs_ocr": False}
        except UnicodeDecodeError:
            continue
    return {"text": "", "pages": 0, "kind": "text", "needs_ocr": False}


# 扩展名 → 解析函数
HANDLERS = {
    ".pdf": parse_pdf,
    ".docx": parse_docx,
    ".xlsx": parse_xlsx,
    ".txt": parse_text,
    ".md": parse_text,
    ".csv": parse_text,
}

SUPPORTED = tuple(HANDLERS)   # 支持的扩展名


def parse_file(path: Path) -> dict:
    """按扩展名分发到对应解析函数，并补齐统一字段。"""
    handler = HANDLERS.get(path.suffix.lower())
    if handler is None:
        return {"text": "", "pages": 0, "kind": path.suffix or "unknown",
                "needs_ocr": False, "uniq_len": 0, "hits": 0,
                "error": "不支持的格式"}
    try:
        r = handler(path)
    except Exception as e:
        return {"text": "", "pages": 0, "kind": path.suffix, "needs_ocr": False,
                "uniq_len": 0, "hits": 0, "error": f"{type(e).__name__}: {e}"}
    # 非 PDF 的解析器没算这两个指标，这里补上
    if "uniq_len" not in r:
        r["uniq_len"], r["hits"] = effective_stats(r["text"])
    r["error"] = ""
    return r


# ============================================================
# 解析预览（命令行）
# ============================================================
def iter_files(root: Path, exts=SUPPORTED):
    for p in sorted(root.rglob("*")):
        if p.is_file() and p.suffix.lower() in exts:
            yield p


def main() -> None:
    ap = argparse.ArgumentParser(description="解析预览：看看资料能抽出多少文本")
    ap.add_argument("--dir", default=str(DATA_ROOT), help="要扫描的目录")
    ap.add_argument("--limit", type=int, default=12, help="最多解析多少个文件")
    args = ap.parse_args()

    root = Path(args.dir)
    if not root.is_dir():
        print(f"目录不存在：{root}")
        return

    files = list(iter_files(root))
    print(f"目录：{root}")
    print(f"可解析文件（{', '.join(SUPPORTED)}）：{len(files)} 个，本次解析前 {args.limit} 个\n")

    head = f"{'文件':<34}{'类型':<6}{'页':>4}{'总字':>7}{'有效':>7}{'词':>4}  判定"
    print(head)
    print("-" * 92)

    ocr_count = 0
    for f in files[: args.limit]:
        r = parse_file(f)
        n = len(r["text"])
        if r["error"]:
            verdict = f"解析失败：{r['error'][:24]}"
        elif r["needs_ocr"]:
            verdict = "⚠️ 需 OCR（抽到的是页眉噪音）"
            ocr_count += 1
        else:
            verdict = "✅ 文本型，可直接用"
        name = f.name if len(f.name) <= 30 else f.name[:27] + "..."
        print(f"{name:<34}{r['kind']:<6}{r['pages']:>4}{n:>7}{r['uniq_len']:>7}{r['hits']:>4}  {verdict}")

    print("-" * 92)
    print("  列说明：总字=抽出字符数；有效=去重后的字符数；词=领域关键词命中数")
    print(f"\n本次样本中，需要 OCR 的：{ocr_count} / {min(args.limit, len(files))}")


if __name__ == "__main__":
    main()
