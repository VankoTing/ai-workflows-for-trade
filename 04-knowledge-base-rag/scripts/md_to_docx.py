"""
Markdown → Word（docx）转换器。

为什么需要它：.md 文件在很多 Windows 上双击打不开（没有关联程序）。
但求职材料需要能直接打开、能编辑、能打印 —— Word 是最稳的。

支持：# ## ### 标题 / > 引用 / - 列表 / 1. 有序列表 / | 表格 |
      **加粗** / `行内代码` / --- 分隔线

用法：
    python scripts/md_to_docx.py "docs/career/项目说明.md"
    python scripts/md_to_docx.py "docs/career/项目说明.md" --out "D:/桌面/项目说明.docx"
"""

from __future__ import annotations

import argparse
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

CN_FONT = "微软雅黑"
MONO = "Consolas"


def _set_cn_font(doc: Document, name: str = CN_FONT) -> None:
    """中文字体必须显式设 eastAsia，否则可能显示成方框。"""
    st = doc.styles["Normal"]
    st.font.name = name
    st.font.size = Pt(10.5)
    rpr = st.element.get_or_add_rPr()
    rf = rpr.get_or_add_rFonts()
    rf.set(qn("w:eastAsia"), name)
    rf.set(qn("w:ascii"), name)
    rf.set(qn("w:hAnsi"), name)


def _runs(p, text: str) -> None:
    """写入带 **加粗** 与 `代码` 的文本。"""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = p.add_run(part[2:-2])
            r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = p.add_run(part[1:-1])
            r.font.name = MONO
            r.font.size = Pt(9.5)
            r.font.color.rgb = RGBColor(0xB0, 0x3A, 0x2E)
        else:
            p.add_run(part)


def _mk_table(doc: Document, rows: list[list[str]]) -> None:
    if not rows:
        return
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=0, cols=cols)
    t.style = "Light Grid Accent 1"
    for i, row in enumerate(rows):
        cells = t.add_row().cells
        for j in range(cols):
            txt = row[j] if j < len(row) else ""
            cells[j].text = ""
            p = cells[j].paragraphs[0]
            _runs(p, txt)
            for r in p.runs:
                r.font.size = Pt(9.5)
                if i == 0:
                    r.bold = True
    doc.add_paragraph()


def convert(md_path: Path, out_path: Path) -> None:
    lines = md_path.read_text(encoding="utf-8").splitlines()
    doc = Document()
    _set_cn_font(doc)

    buf_tbl: list[list[str]] = []

    def flush() -> None:
        nonlocal buf_tbl
        if buf_tbl:
            _mk_table(doc, buf_tbl)
            buf_tbl = []

    for raw in lines:
        line = raw.rstrip()

        if line.strip().startswith("|") and line.count("|") >= 2:
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c or "-") for c in cells):
                continue                       # 表格分隔行
            buf_tbl.append(cells)
            continue
        flush()

        if not line.strip():
            continue
        if re.fullmatch(r"-{3,}", line.strip()):
            doc.add_paragraph("─" * 42).alignment = WD_ALIGN_PARAGRAPH.CENTER
            continue
        if line.startswith("### "):
            doc.add_heading(line[4:].strip(), level=3)
            continue
        if line.startswith("## "):
            doc.add_heading(line[3:].strip(), level=2)
            continue
        if line.startswith("# "):
            doc.add_heading(line[2:].strip(), level=1)
            continue
        if line.startswith("> "):
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Pt(18)
            _runs(p, line[2:])
            for r in p.runs:
                r.italic = True
                r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            continue
        m = re.match(r"^(\s*)([-*])\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.left_indent = Pt(18 + len(m.group(1)) * 6)
            _runs(p, m.group(3))
            continue
        m = re.match(r"^(\s*)(\d+)\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.left_indent = Pt(18 + len(m.group(1)) * 6)
            _runs(p, m.group(3))
            continue
        p = doc.add_paragraph()
        _runs(p, line.strip())

    flush()
    out_path.parent.mkdir(parents=True, exist_ok=True)
    doc.save(out_path)
    print(f"已生成：{out_path}")


def main() -> None:
    ap = argparse.ArgumentParser()
    ap.add_argument("md", help="Markdown 文件路径")
    ap.add_argument("--out", default=None, help="输出 docx 路径（默认同目录同名）")
    args = ap.parse_args()

    p = Path(args.md)
    if not p.is_absolute():
        p = Path(__file__).resolve().parents[1] / p
    if not p.exists():
        print(f"找不到文件：{p}")
        return
    out = Path(args.out) if args.out else p.with_suffix(".docx")
    convert(p, out)


if __name__ == "__main__":
    main()
