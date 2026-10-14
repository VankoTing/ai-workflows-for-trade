"""往返测试：读参照文件 -> 转成 InquiryTable -> 重新导出 -> 逐项对比。

用来验证 exporter 是否精确复刻了参照格式。
用法：python tools/roundtrip_test.py <参照文件.xlsx>
"""

from __future__ import annotations

import re
import sys
from pathlib import Path

PROJECT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PROJECT / "src"))

from openpyxl import load_workbook  # noqa: E402

from exporter import BiCell, InquiryRow, InquiryTable, export_inquiry  # noqa: E402

CJK = re.compile(r"[\u4e00-\u9fff]")
TITLE_RE = re.compile(r"Ref\.\s*(.+?)\s*\((.+)\)\s*$")
NOTE_RE = re.compile(r"^Note\s*\d+\s*[:：]\s*")
NOTE_ZH_RE = re.compile(r"^备注\s*\d+\s*[:：]\s*")


def split_bi(value) -> BiCell:
    text = "" if value is None else str(value)
    if "\n" in text:
        en, zh = text.split("\n", 1)
        return BiCell(en=en.strip(), zh=zh.strip())
    text = text.strip()
    if not text:
        return BiCell()
    return BiCell(zh=text) if CJK.search(text) else BiCell(en=text)


def read_table(path: Path) -> InquiryTable:
    wb = load_workbook(path, data_only=True)
    ws = wb.worksheets[0]
    table = InquiryTable()

    match = TITLE_RE.search(str(ws["A1"].value or ""))
    if match:
        table.ref = match.group(1).strip()
        table.country = match.group(2).strip()

    total_row = None
    for r in range(3, ws.max_row + 1):
        if str(ws.cell(row=r, column=1).value or "").startswith("Total"):
            total_row = r
            break
    last_data = (total_row - 1) if total_row else ws.max_row

    for r in range(3, last_data + 1):
        if not ws.cell(row=r, column=2).value:
            continue
        table.rows.append(
            InquiryRow(
                no=str(ws.cell(row=r, column=1).value or ""),
                valve_type=split_bi(ws.cell(row=r, column=2).value),
                size=split_bi(ws.cell(row=r, column=3).value),
                rating=split_bi(ws.cell(row=r, column=4).value),
                end_connection=split_bi(ws.cell(row=r, column=5).value),
                port=split_bi(ws.cell(row=r, column=6).value),
                body_material=split_bi(ws.cell(row=r, column=7).value),
                trim=split_bi(ws.cell(row=r, column=8).value),
                design_standard=split_bi(ws.cell(row=r, column=9).value),
                special_requirements=split_bi(ws.cell(row=r, column=10).value),
                qty=str(ws.cell(row=r, column=11).value or ""),
                unit=split_bi(ws.cell(row=r, column=12).value),
            )
        )

    start = (total_row + 2) if total_row else (last_data + 2)
    for r in range(start + 1, ws.max_row + 1):
        raw = ws.cell(row=r, column=1).value
        if not raw:
            continue
        text = NOTE_RE.sub("", str(raw))
        if "\n" in text:
            en, zh = text.split("\n", 1)
            table.notes.append(BiCell(en=en.strip(), zh=NOTE_ZH_RE.sub("", zh.strip())))
        else:
            table.notes.append(BiCell(en=text.strip()))
    wb.close()
    return table


def norm_content(value) -> str:
    """把「两行内容完全相同」归一化成一行。

    exporter 会自动合并重复的双语行（原文件里存在 "5000 PSI\\n5000 PSI" 这种），
    这是有意的改进，对比时视为等价。
    """
    text = "" if value is None else str(value)
    parts = text.split("\n")
    if len(parts) == 2 and parts[0].strip() == parts[1].strip():
        return parts[0].strip()
    return text


def cell_sig(cell) -> str:
    bits: list[str] = []
    font = cell.font
    if font:
        bits.append(f"bold{int(bool(font.bold))}")
        bits.append(f"sz{font.size}")
        try:
            if font.color and isinstance(font.color.rgb, str):
                bits.append(f"fg{font.color.rgb}")
        except Exception:
            pass
    try:
        fill = cell.fill
        if fill and fill.fill_type == "solid" and isinstance(fill.start_color.rgb, str):
            bits.append(f"bg{fill.start_color.rgb}")
    except Exception:
        pass
    al = cell.alignment
    if al:
        bits.append(f"h{al.horizontal}")
        bits.append(f"v{al.vertical}")
        bits.append(f"wrap{int(bool(al.wrap_text))}")
    bd = cell.border
    if bd:
        for name in ("left", "right", "top", "bottom"):
            side = getattr(bd, name, None)
            bits.append(f"{name[:1]}{side.style if side and side.style else '-'}")
    return " ".join(bits)


def compare(orig: Path, new: Path) -> list[str]:
    wa = load_workbook(orig, data_only=True).worksheets[0]
    wb_ = load_workbook(new, data_only=True).worksheets[0]
    diffs: list[str] = []

    if wa.title != wb_.title:
        diffs.append(f"sheet 名不同：{wa.title!r} vs {wb_.title!r}")

    ma = sorted(str(x) for x in wa.merged_cells.ranges)
    mb = sorted(str(x) for x in wb_.merged_cells.ranges)
    if ma != mb:
        diffs.append(f"合并单元格不同：\n      原={ma}\n      新={mb}")

    for col in "ABCDEFGHIJKL":
        va = wa.column_dimensions[col].width
        vb = wb_.column_dimensions[col].width
        same = (va is None and vb is None) or (
            va is not None and vb is not None and abs(va - vb) < 0.05
        )
        if not same:
            diffs.append(f"列宽 {col} 不同：{va} vs {vb}")

    # 合并区域里除左上角以外的单元格不参与显示，跳过它们的样式对比
    skip: set[tuple[int, int]] = set()
    for rng in wa.merged_cells.ranges:
        for rr in range(rng.min_row, rng.max_row + 1):
            for cc in range(rng.min_col, rng.max_col + 1):
                if (rr, cc) != (rng.min_row, rng.min_col):
                    skip.add((rr, cc))

    max_row = max(wa.max_row, wb_.max_row)
    for r in range(1, max_row + 1):
        ha = wa.row_dimensions[r].height
        hb = wb_.row_dimensions[r].height
        same = (ha is None and hb is None) or (
            ha is not None and hb is not None and abs(ha - hb) < 0.05
        )
        if not same:
            diffs.append(f"行高 第{r}行 不同：{ha} vs {hb}")

        # 整行两边都没有内容时，忽略样式差异（空行的单元格样式不影响显示）
        row_empty = all(
            not str(wa.cell(row=r, column=c).value or "").strip()
            and not str(wb_.cell(row=r, column=c).value or "").strip()
            for c in range(1, 13)
        )

        for c in range(1, 13):
            if (r, c) in skip:
                continue
            ca = wa.cell(row=r, column=c)
            cb = wb_.cell(row=r, column=c)
            va = norm_content(ca.value)
            vb = norm_content(cb.value)
            if va != vb:
                diffs.append(
                    f"内容 {ca.coordinate} 不同：\n      原={str(va)[:80]!r}\n      新={str(vb)[:80]!r}"
                )
            sa, sb = cell_sig(ca), cell_sig(cb)
            if sa != sb and not row_empty:
                diffs.append(f"样式 {ca.coordinate} 不同：\n      原={sa}\n      新={sb}")
    return diffs


def main() -> None:
    if len(sys.argv) < 2:
        print("用法：python tools/roundtrip_test.py <参照文件.xlsx>")
        return
    orig = Path(sys.argv[1])
    table = read_table(orig)
    print(f"从参照文件解析出：ref={table.ref!r} country={table.country!r}")
    print(f"  数据行 {len(table.rows)} 条，备注 {len(table.notes)} 条，合计数量 {table.total_qty()}")

    out = PROJECT / "output" / f"roundtrip_{orig.stem}.xlsx"
    export_inquiry(table, out)
    print(f"  已重新导出：{out.name}")

    # 顺便存一份抽取结果 JSON，用来测试 --from-json 离线导出链路
    import json

    from extractor import FIELDS

    payload = {
        "ref": table.ref,
        "country": table.country,
        "rows": [
            {
                "no": row.no,
                "qty": row.qty,
                **{
                    f: {"en": getattr(row, f).en, "zh": getattr(row, f).zh}
                    for f in FIELDS
                },
            }
            for row in table.rows
        ],
        "notes": [{"en": n.en, "zh": n.zh} for n in table.notes],
    }
    json_path = PROJECT / "output" / f"{orig.stem}_extracted.json"
    json_path.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"  已导出抽取 JSON：{json_path.name}")

    diffs = compare(orig, out)
    print()
    if not diffs:
        print("✅ 完全一致：sheet 名 / 合并单元格 / 列宽 / 行高 / 内容 / 样式 全部吻合")
    else:
        print(f"❌ 发现 {len(diffs)} 处差异：")
        for d in diffs[:60]:
            print("   -", d)
        if len(diffs) > 60:
            print(f"   … 还有 {len(diffs) - 60} 处")


if __name__ == "__main__":
    main()
