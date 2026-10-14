"""读取一个 xlsx 的完整结构（内容 + 格式），用于复刻模板。

用法：python tools/inspect_template.py <xlsx路径> [--brief]

输出：子表列表、每个子表的行内容、合并单元格、列宽行高、冻结窗格，以及单元格样式标记。
--brief 只输出内容，不输出样式标记。
"""

from __future__ import annotations

import sys
from pathlib import Path

from openpyxl import load_workbook


def fmt_style(cell) -> str:
    bits: list[str] = []
    font = cell.font
    if font:
        if font.bold:
            bits.append("bold")
        if font.size and abs(font.size - 11) > 0.01:
            bits.append(f"sz{font.size}")
        try:
            if font.color and font.color.rgb and isinstance(font.color.rgb, str):
                bits.append(f"fg{font.color.rgb}")
        except Exception:
            pass
    try:
        fill = cell.fill
        if fill and fill.fill_type == "solid" and fill.start_color and isinstance(fill.start_color.rgb, str):
            bits.append(f"bg{fill.start_color.rgb}")
    except Exception:
        pass
    al = cell.alignment
    if al:
        if al.wrap_text:
            bits.append("wrap")
        if al.horizontal:
            bits.append(al.horizontal)
    if cell.number_format and cell.number_format != "General":
        bits.append(f"nf:{cell.number_format}")
    bd = cell.border
    if bd:
        sides = []
        for name in ("left", "right", "top", "bottom"):
            side = getattr(bd, name, None)
            if side is not None and side.style:
                sides.append(f"{name[0]}={side.style}")
        if sides:
            bits.append("bd[" + ",".join(sides) + "]")
    return " ".join(bits)


def main() -> None:
    if len(sys.argv) < 2:
        print("用法：python tools/inspect_template.py <xlsx路径> [--brief]")
        return
    path = Path(sys.argv[1])
    brief = "--brief" in sys.argv

    wb = load_workbook(path, data_only=True)
    print(f"文件：{path.name}")
    print(f"子表顺序：{[ws.title for ws in wb.worksheets]}")

    for ws in wb.worksheets:
        print()
        print("=" * 78)
        print(f"## 子表：{ws.title}   尺寸 {ws.max_row} 行 x {ws.max_column} 列")
        print(f"   冻结窗格：{ws.freeze_panes}")
        merged = sorted(str(r) for r in ws.merged_cells.ranges)
        print(f"   合并单元格：{merged if merged else '无'}")
        widths = {
            k: round(v.width, 1)
            for k, v in ws.column_dimensions.items()
            if getattr(v, "width", None)
        }
        print(f"   列宽：{widths}")
        heights = {
            k: v.height for k, v in ws.row_dimensions.items() if getattr(v, "height", None)
        }
        if heights:
            print(f"   行高：{heights}")
        print()

        for row in ws.iter_rows():
            for cell in row:
                if cell.value is None or str(cell.value).strip() == "":
                    continue
                text = str(cell.value).replace("\n", "\\n")
                if len(text) > 150:
                    text = text[:150] + "…"
                style = "" if brief else fmt_style(cell)
                tag = f"   [{style}]" if style else ""
                print(f"  {cell.coordinate:<6}| {text}{tag}")
    wb.close()


if __name__ == "__main__":
    main()
