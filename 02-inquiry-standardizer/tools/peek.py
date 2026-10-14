"""快速查看生成的 xlsx 内容。

用法：python tools/peek.py output/xxx.xlsx
"""

import sys
from pathlib import Path

from openpyxl import load_workbook


def main() -> None:
    path = Path(sys.argv[1])
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        print(f"=== {ws.title}  ({ws.max_row} 行 x {ws.max_column} 列)")
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            cells = ["" if c is None else str(c).replace("\n", " ") for c in row]
            while cells and not cells[-1].strip():
                cells.pop()
            if not cells:
                print(f"{i:>3}|")
                continue
            print(f"{i:>3}| " + " | ".join(cells)[:300])
    wb.close()


if __name__ == "__main__":
    main()
