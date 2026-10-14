"""MTO 中英对照：保留原表结构，给列名和单元格值补中文。

策略是「按列提取唯一值 -> 批量翻译 -> 本地回填」：
388 行里真正不同的取值只有一百多个，翻译一次全程复用。
token 消耗极小，而且同一个值在任何位置译法都一致。

列名单独处理：带上该列的样本值一起翻译，避免同一个词在不同列被译错
（例如 UNIT 在「单元编号」列应译「单元」，在「数量单位」列应译「单位」）。

用法：
    python src/mto.py <脱敏后的mto.xlsx> [-o 输出.xlsx]
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import re
from datetime import datetime
from pathlib import Path

import httpx
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

try:
    from dotenv import load_dotenv

    load_dotenv()
except ImportError:  # pragma: no cover
    pass

PROJECT = Path(__file__).resolve().parent.parent
OUTPUT_DIR = PROJECT / "output"

DATA_START = 7
HEADER_ROW_MAIN = 5
HEADER_ROW_SUB = 6
HEADER_LAST = 6

# 脱敏后的代号列，本身没有语义，不翻译
SKIP_TRANSLATE_COLS = {2, 5, 6}   # B 阀门位号 / E P&ID / F 管线号

# 居中对齐的列（编号、尺寸、等级、数量这类短字段）
CENTER_COLS = {1, 3, 12, 13, 14, 17, 20, 21, 22, 25, 26, 27}

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD_FILL = PatternFill("solid", fgColor="FF1F4E79")
# 动态补充列用橙色表头，一眼能和原表列区分开
EXTRA_FILL = PatternFill("solid", fgColor="FFC55A11")
HEAD_FONT = Font(size=10, bold=True, color="FFFFFFFF")
BODY_FONT = Font(size=10)
WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

DEFAULT_BASE_URL = "https://api.deepseek.com"
DEFAULT_MODEL = "deepseek-chat"

SYSTEM_PROMPT = """你是资深工业阀门与电厂管道专业的翻译。

你会收到两部分内容：
  columns: 表格的列名，每项含 id、原文 text、以及该列的样本值 samples
  values : 需要翻译的零散字符串

请分别翻译后，按指定结构返回。

【术语规范（来自电厂阀门材料清单语境）】
- 阀门类型：GATE→闸阀、GLOBE→截止阀、CHECK→止回阀、BALL→球阀、BUTTERFLY→蝶阀、
  NEEDLE→针型阀、PLUG→旋塞阀、DIAPHRAGM→隔膜阀、SWING CHECK→旋启式止回阀
- 连接方式：BW→对焊、SW→承插焊、Socket Welded→承插焊、Butt Welded→对焊、
  Flanged→法兰连接、RF→凸面法兰、RTJ→环形槽面、NPT→NPT 螺纹
- 操作方式：Manual→手动、M→手动、Motor Operated→电动、Pneumatic→气动、Gear→齿轮传动
- 压力温度列：P kgf/cm2→压力 (kgf/cm²)、Т,ºс→温度 (℃)
- 其他：pcs/PCS→个、set→套、Spare→备品、Q'ty→数量、Rev.→版本、Remarks→备注

【列名翻译规则】
- 必须结合 samples 判断该列的真实含义，同一个词在不同列可以有不同译法。
  例：UNIT 在值形如 UNIT#1 的列译「单元」，在值形如 PCS 的列译「单位」。
- 图纸号/编号类列名同样要翻译，不要因为它是缩写或代号就跳过：
  P&ID→P&ID 图号、VALVE TAG NUMBER→阀门位号、PIPE LINE/EQP NO.→管道/设备编号、
  Piping Material Specification Code→管道材料等级代码、Q'ty Spare→备品数量
- 上级表头 + 下级表头要合成一个通顺的列名。
  例：上级 Medium operating parameters + 本级 P kgf/cm2 → 介质操作压力
      上级 Medium operating parameters + 本级 Т,ºс       → 介质操作温度
      上级 Medium Design Pressure  + 本级 P kgf/cm2      → 介质设计压力
      上级 Medium Design Pressure  + 本级 Т,ºс           → 介质设计温度

【必须遵守】
- 标准号、材质牌号、型号代码一律**保持原文**，可在其后用中文括号补充含义。
  例：ASTM A216 GR.WCB → ASTM A216 GR.WCB（碳钢铸件）
      ASME B16.34/SA182-F91 → ASME B16.34/SA182-F91（F91 合金钢锻件）
      A-182 F91 → A-182 F91（F91 合金钢）
- 代号（形如 REF-0001）原样保留，绝不翻译、不改写。
- 纯数字、纯符号、单位缩写（kgf/cm2、ºс 等）原样保留。
- 已经是中文的原样返回。译文简洁专业，不要加解释。

【输出格式】只输出 JSON，不要其他文字：
{
  "columns": {"<列 id>": "中文列名"},
  "values": {"<原文>": "中文译文"}
}"""


REVIEW_PROMPT = """你是资深工业阀门外贸技术支持工程师。

下面给你一份电厂阀门材料清单（MTO）的**列结构**和每列的样本值。
请审阅：**从报价的角度看，这份清单还缺少哪些必要的技术参数？**

【判定标准】三条必须同时满足，才把该参数列进结果：
1. 属于阀门报价必需、或会严重影响报价准确度
2. 清单现有的列**没有覆盖**它（列名和样本值都看一遍再判断）
3. 整份清单都没有提供这项信息

【常见候选，只在确实缺失时才加】
Trim (Ball/Stem/Seat) / 内件材质、Design Standard / 设计标准、
Face-to-Face (FF) / 结构长度、Flange Standard / 法兰标准、
Face Finish / 密封面粗糙度、Painting / 涂装要求、
Inspection & Test / 检验与试验、Third-party Certification / 第三方认证、
Applicable Temperature Range / 适用温度范围、Test Pressure / 试验压力

【硬性要求】
- 最多 5 列，宁缺毋滥。确实不缺就返回空数组 []。
- key 用小写蛇形英文，如 trim、design_standard。
- 不要重复清单里已经有的列（例如已经有 BODY MATERIAL，就不要再提 Body Material）。

【输出】只输出 JSON，不要其他文字：
{
  "extra_columns": [
    {"key": "trim", "en": "Trim (Ball/Stem/Seat)", "zh": "内件材质"}
  ],
  "reason": "一句话说明这些为什么是报价必需的"
}"""


def read_sheet_structure(path: Path):
    """读出每张表：列名项（含样本）、列宽、数据行。"""
    wb = load_workbook(path, data_only=True)
    sheets = []
    for ws in wb.worksheets:
        # 先把合并单元格展开成 坐标 -> 左上角值
        merged: dict[tuple[int, int], object] = {}
        for rng in ws.merged_cells.ranges:
            top_left = ws.cell(row=rng.min_row, column=rng.min_col).value
            for rr in range(rng.min_row, rng.max_row + 1):
                for cc in range(rng.min_col, rng.max_col + 1):
                    merged[(rr, cc)] = top_left

        def hv(row: int, col: int):
            if (row, col) in merged:
                return merged[(row, col)]
            return ws.cell(row=row, column=col).value

        ncols = ws.max_column
        rows = []
        for r in range(DATA_START, ws.max_row + 1):
            if ws.cell(row=r, column=2).value in (None, ""):
                continue
            rows.append([ws.cell(row=r, column=c).value for c in range(1, ncols + 1)])

        columns = []
        for c in range(1, ncols + 1):
            main = str(hv(HEADER_ROW_MAIN, c) or "").replace("\n", " ").strip()
            sub = str(hv(HEADER_ROW_SUB, c) or "").replace("\n", " ").strip()
            if sub and sub != main:
                key = f"{main} / {sub}"
                en = f"{main}\n{sub}" if main else sub
            else:
                key = main or sub
                en = main or sub
            samples = []
            if c not in SKIP_TRANSLATE_COLS:
                for row in rows:
                    v = row[c - 1]
                    if v not in (None, ""):
                        s = str(v).strip()
                        if s not in samples:
                            samples.append(s)
                    if len(samples) >= 4:
                        break
            columns.append({"id": get_column_letter(c), "key": key, "en": en, "samples": samples})

        widths = {}
        for c in range(1, ncols + 1):
            letter = get_column_letter(c)
            w = ws.column_dimensions[letter].width if letter in ws.column_dimensions else None
            if w:
                widths[c] = w

        sheets.append({"title": ws.title, "columns": columns, "rows": rows,
                       "widths": widths, "ncols": ncols})
    wb.close()
    return sheets


def collect_strings(sheets):
    """收集待翻译内容：列名项 + 零散值。"""
    col_items: list[dict] = []
    seen_ids: set[str] = set()
    values: dict[str, None] = {}
    for sh in sheets:
        # 列名一律翻译。SKIP_TRANSLATE_COLS 只表示「这一列的值是代号、不用翻」，
        # 不代表列名也不用翻 —— VALVE TAG NUMBER / P&ID 这些列名本身是有语义的。
        # 按列 id 收集，不按词去重：同一个词在不同列可能就是不同译法（UNIT 单元 / UNIT 单位）
        for col in sh["columns"]:
            if col["id"] in seen_ids:
                continue
            if col["key"] and re.search(r"[A-Za-z]", col["key"]):
                seen_ids.add(col["id"])
                col_items.append({"id": col["id"], "key": col["key"],
                                  "text": col["key"], "samples": col["samples"][:4]})
        for row in sh["rows"]:
            for c, v in enumerate(row, start=1):
                if c in SKIP_TRANSLATE_COLS or v in (None, ""):
                    continue
                t = str(v).strip()
                if t and re.search(r"[A-Za-z]", t):
                    values.setdefault(t, None)
    return col_items, list(values)


async def translate_batch(col_items, values, api_key, model, base_url):
    user = json.dumps(
        {
            "columns": [{"id": it["id"], "text": it["text"], "samples": it["samples"]} for it in col_items],
            "values": values,
        },
        ensure_ascii=False,
    )
    body = {
        "model": model,
        "messages": [
            {"role": "system", "content": SYSTEM_PROMPT},
            {"role": "user", "content": user},
        ],
        "temperature": 0.1,
        "response_format": {"type": "json_object"},
        "stream": False,
    }
    async with httpx.AsyncClient(timeout=300.0) as client:
        resp = await client.post(
            f"{base_url.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json=body,
        )
        resp.raise_for_status()
        content = resp.json()["choices"][0]["message"]["content"]

    text = content.strip()
    fence = re.search(r"```(?:json)?\s*(.+?)\s*```", text, re.DOTALL)
    if fence:
        text = fence.group(1)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        s, e = text.find("{"), text.rfind("}")
        data = json.loads(text[s : e + 1]) if s >= 0 and s < e else {}

    cols = {str(k).strip(): str(v).strip() for k, v in (data.get("columns") or {}).items()}
    vals = {str(k).strip(): str(v).strip() for k, v in (data.get("values") or {}).items()}
    return cols, vals


async def review_missing(sheets, api_key, model, base_url):
    """审阅 MTO，找出清单里缺失的必要参数，用于在表头追加列。"""
    brief = []
    for sh in sheets:
        cols = [
            {"col": c["id"], "name": c["key"], "samples": c["samples"][:3]}
            for c in sh["columns"]
            if c["key"]
        ]
        brief.append({"sheet": sh["title"], "columns": cols})

    body = {
        "model": model,
        "messages": [
            {"role": "system", "content": REVIEW_PROMPT},
            {"role": "user", "content": json.dumps(brief, ensure_ascii=False)},
        ],
        "temperature": 0.1,
        "response_format": {"type": "json_object"},
        "stream": False,
    }
    async with httpx.AsyncClient(timeout=300.0) as client:
        resp = await client.post(
            f"{base_url.rstrip('/')}/chat/completions",
            headers={"Authorization": f"Bearer {api_key}", "Content-Type": "application/json"},
            json=body,
        )
        resp.raise_for_status()
        content = resp.json()["choices"][0]["message"]["content"]

    text = content.strip()
    fence = re.search(r"```(?:json)?\s*(.+?)\s*```", text, re.DOTALL)
    if fence:
        text = fence.group(1)
    try:
        data = json.loads(text)
    except json.JSONDecodeError:
        s, e = text.find("{"), text.rfind("}")
        data = json.loads(text[s : e + 1]) if s >= 0 and s < e else {}

    columns = []
    for item in (data.get("extra_columns") or [])[:5]:
        if not isinstance(item, dict):
            continue
        key = str(item.get("key", "") or "").strip()
        en = str(item.get("en", "") or "").strip()
        zh = str(item.get("zh", "") or "").strip()
        if key and (en or zh):
            columns.append({"key": key, "en": en, "zh": zh})
    return columns, str(data.get("reason", "") or "").strip()


MISSING_TEXT = "Not provided\n未提供"


def export(sheets, col_tr: dict, val_tr: dict, out_path: Path, title: str,
           extra_columns: list[dict] | None = None) -> Path:
    wb = Workbook()
    first = True
    for sh in sheets:
        ws = wb.active if first else wb.create_sheet()
        first = False
        ws.title = sh["title"][:31]
        extras = extra_columns or []
        base_cols = sh["ncols"]
        ncols = base_cols + len(extras)

        ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
        cell = ws.cell(row=1, column=1, value=f"{title} — {sh['title']}")
        cell.font = Font(size=12, bold=True, color="FFFFFFFF")
        cell.fill = HEAD_FILL
        cell.alignment = WRAP_CENTER
        ws.row_dimensions[1].height = 22

        for c in range(1, base_cols + 1):
            col = sh["columns"][c - 1]
            zh = col_tr.get(col["id"], "") or col_tr.get(col["key"], "")
            if not col["en"]:
                value = ""          # 原表里没内容的空列，保持空白，不要填占位符
            elif zh and zh != col["key"]:
                value = f"{col['en']}\n{zh}"
            else:
                value = col["en"]
            hc = ws.cell(row=2, column=c, value=value)
            hc.font = HEAD_FONT
            hc.fill = HEAD_FILL
            hc.alignment = WRAP_CENTER
            hc.border = BORDER

        # 动态补充列：AI 审阅后判定「这份清单缺少的必要参数」
        for j, extra in enumerate(extras):
            hc = ws.cell(
                row=2,
                column=base_cols + j + 1,
                value=f"{extra['en']}\n{extra['zh']}" if extra.get("zh") else extra.get("en", ""),
            )
            hc.font = HEAD_FONT
            hc.fill = EXTRA_FILL
            hc.alignment = WRAP_CENTER
            hc.border = BORDER
        ws.row_dimensions[2].height = 38

        for i, row in enumerate(sh["rows"], start=3):
            for c in range(1, base_cols + 1):
                v = row[c - 1]
                if c in SKIP_TRANSLATE_COLS or v in (None, ""):
                    rendered = "" if v in (None, "") else str(v)
                else:
                    raw = str(v).strip()
                    zh = val_tr.get(raw)
                    rendered = f"{raw}\n{zh}" if zh and zh != raw else raw
                dc = ws.cell(row=i, column=c, value=rendered)
                dc.font = BODY_FONT
                dc.alignment = WRAP_CENTER if c in CENTER_COLS else WRAP_LEFT
                dc.border = BORDER

            # 补充列：清单里没有的必填参数，逐行标「未提供」
            for j in range(len(extras)):
                dc = ws.cell(row=i, column=base_cols + j + 1, value=MISSING_TEXT)
                dc.font = BODY_FONT
                dc.alignment = WRAP_CENTER
                dc.border = BORDER
            ws.row_dimensions[i].height = 30

        for c in range(1, base_cols + 1):
            w = sh["widths"].get(c)
            if w:
                ws.column_dimensions[get_column_letter(c)].width = min(max(w, 8), 34)
        for j in range(len(extras)):
            ws.column_dimensions[get_column_letter(base_cols + j + 1)].width = 20
        ws.freeze_panes = "A3"

    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path


async def main_async(args) -> None:
    src = Path(args.source)
    if not src.exists():
        print(f"文件不存在：{src}")
        return

    sheets = read_sheet_structure(src)
    total = sum(len(s["rows"]) for s in sheets)
    print(f"读入 {len(sheets)} 张表，共 {total} 行")
    for s in sheets:
        print(f"  {s['title']}: {len(s['rows'])} 行 x {s['ncols']} 列")

    col_items, values = collect_strings(sheets)
    print(f"\n待翻译：列名 {len(col_items)} 个，取值 {len(values)} 个（已去重）")

    api_key = os.getenv("DEEPSEEK_API_KEY", "").strip()
    if not api_key:
        print("! 未配置 DEEPSEEK_API_KEY")
        return
    model = os.getenv("DEEPSEEK_MODEL", DEFAULT_MODEL).strip()
    base_url = os.getenv("DEEPSEEK_BASE_URL", DEFAULT_BASE_URL).strip()

    print(f"正在调用 {model}…")
    col_tr, val_tr = await translate_batch(col_items, values, api_key, model, base_url)
    print(f"  返回：列名 {len(col_tr)} 条，取值 {len(val_tr)} 条")

    missing = [v for v in values if v not in val_tr]
    if missing:
        print(f"  ! {len(missing)} 条取值未返回，保持原文：{missing[:4]}")

    # 输出前审阅：这份清单还缺哪些必要参数
    print("\n正在审阅清单缺失的必要参数…")
    extra_columns, reason = await review_missing(sheets, api_key, model, base_url)
    if extra_columns:
        print(f"  判定缺少 {len(extra_columns)} 项必要参数，追加为补充列（橙色表头）：")
        for col in extra_columns:
            print(f"     - {col['en']} / {col['zh']}")
        if reason:
            print(f"    理由：{reason}")
    else:
        print("  未发现缺失的必要参数，不追加列。")

    title = args.title or f"{src.stem} 中英对照"
    out = Path(args.out) if args.out else OUTPUT_DIR / f"{src.stem}_Bilingual_{datetime.now():%Y%m%d_%H%M}.xlsx"
    export(sheets, col_tr, val_tr, out, title, extra_columns)

    cache = out.with_name(out.stem + "_translation.json")
    cache.write_text(
        json.dumps({"columns": col_tr, "values": val_tr}, ensure_ascii=False, indent=2),
        encoding="utf-8",
    )
    print(f"\n已导出：{out}")
    print(f"译文对照：{cache.name}")


def main() -> None:
    parser = argparse.ArgumentParser(description="MTO 中英对照")
    parser.add_argument("source", help="脱敏后的 xlsx")
    parser.add_argument("-o", "--out", default=None, help="输出 xlsx")
    parser.add_argument("--title", default=None, help="表头标题")
    args = parser.parse_args()
    asyncio.run(main_async(args))


if __name__ == "__main__":
    main()
