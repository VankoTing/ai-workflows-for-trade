"""导出中英对照询价表 —— 严格复刻既定交付格式。

格式规格来自参照文件 691-2610-26_Inquiry_Bilingual.xlsx：

    第 1 行   A1:L1 合并标题，深蓝底 1F4E79 / 白字 14 加粗 / 居中换行 / 行高 22.5
    第 2 行   表头 12 列，中蓝底 2E75B6 / 白字 10 加粗 / 居中换行 / 四边细框 / 行高 33
    第 3+ 行  数据行，字 10 换行 / 四边细框 / 行高 48
              A-F 与 K-L 居中，G-J 左对齐
    合计行    A=Total/合计、K=总数量、L=单位，浅蓝底 DDEBF7 / 加粗 / 行高 19.5
    空行
    备注标题  A:L 合并，深蓝底白字加粗 / 行高 19.5
    备注行    每条 A:L 合并，字 10 换行左对齐 / 行高 34.5

所有双语单元格内容一律「英文在上、中文在下」两行。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

# --------------------------------------------------------------------------
# 列定义：(英文列名, 中文列名, 列宽, 水平对齐)
# --------------------------------------------------------------------------

BASE_COLUMNS: list[tuple[str, str, float | None, str]] = [
    ("No.", "序号", 5.7, "center"),
    ("Valve Type", "阀门类型", 12.7, "center"),
    ("Size", "尺寸", 13.7, "center"),
    ("Rating", "压力等级", 14.7, "center"),
    ("End Connection", "端部连接", 15.7, "center"),
    ("Port", "通径型式", 12.7, "center"),
    ("Body Material", "阀体材质", 19.7, "left"),
    ("Trim (Ball/Stem/Seat)", "内件（球体/阀杆/阀座）", 26.7, "left"),
    ("Design Standard", "设计标准", 24.7, "left"),
    ("Special Requirements", "特殊要求", 33.7, "left"),
    ("Qty", "数量", 6.7, "center"),
    ("Unit", "单位", None, "center"),
]

THIN = Side(style="thin")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

TITLE_FILL = PatternFill("solid", fgColor="FF1F4E79")
HEAD_FILL = PatternFill("solid", fgColor="FF2E75B6")
TOTAL_FILL = PatternFill("solid", fgColor="FFDDEBF7")

TITLE_FONT = Font(size=14, bold=True, color="FFFFFFFF")
HEAD_FONT = Font(size=10, bold=True, color="FFFFFFFF")
NOTE_TITLE_FONT = Font(size=11, bold=True, color="FFFFFFFF")
BODY_FONT = Font(size=10)
TOTAL_FONT = Font(size=10, bold=True)

WRAP_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
WRAP_LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True)

ROW_H_TITLE = 22.5
ROW_H_HEAD = 33.0
ROW_H_BODY = 48.0
ROW_H_TOTAL = 19.5
ROW_H_NOTE_TITLE = 19.5
ROW_H_NOTE = 34.5


def _numeric(value: object):
    """能转成数字就转，保持 Excel 里是数值而不是文本。"""
    text = str(value).strip()
    if not text:
        return ""
    try:
        return int(text)
    except ValueError:
        pass
    try:
        return float(text)
    except ValueError:
        return text


@dataclass
class BiCell:
    """一个中英双语单元格。"""

    en: str = ""
    zh: str = ""

    def render(self) -> str:
        en, zh = (self.en or "").strip(), (self.zh or "").strip()
        # 两行内容相同时只保留一行，避免出现 "Class 150 / Class 150" 这种重复
        if en and zh and en.casefold() != zh.casefold():
            return f"{en}\n{zh}"
        return en or zh or "-"

    @classmethod
    def of(cls, en: str = "", zh: str = "") -> "BiCell":
        return cls(en=en, zh=zh)


@dataclass
class ExtraColumn:
    """动态补充列：由 AI 审阅后判定「这份询价单缺的必要参数」。

    例如整份单子都没写工作介质，就补一列 Working Fluid / 工作介质，
    对应产品行里填 Not provided / 未提供。
    """

    key: str = ""
    en: str = ""
    zh: str = ""
    width: float = 20.0
    align: str = "left"

    def header(self) -> str:
        return f"{self.en}\n{self.zh}" if self.en and self.zh else (self.en or self.zh)


@dataclass
class InquiryRow:
    no: str = ""
    valve_type: BiCell = field(default_factory=BiCell)
    size: BiCell = field(default_factory=BiCell)
    rating: BiCell = field(default_factory=BiCell)
    end_connection: BiCell = field(default_factory=BiCell)
    port: BiCell = field(default_factory=BiCell)
    body_material: BiCell = field(default_factory=BiCell)
    trim: BiCell = field(default_factory=BiCell)
    design_standard: BiCell = field(default_factory=BiCell)
    special_requirements: BiCell = field(default_factory=BiCell)
    qty: str = ""
    unit: BiCell = field(default_factory=BiCell)
    extra: dict[str, BiCell] = field(default_factory=dict)

    def to_cells(self) -> list:
        return [
            _numeric(self.no),
            self.valve_type.render(),
            self.size.render(),
            self.rating.render(),
            self.end_connection.render(),
            self.port.render(),
            self.body_material.render(),
            self.trim.render(),
            self.design_standard.render(),
            self.special_requirements.render(),
            _numeric(self.qty),
            self.unit.render(),
        ]


@dataclass
class InquiryTable:
    ref: str = ""
    country: str = ""
    rows: list[InquiryRow] = field(default_factory=list)
    notes: list[BiCell] = field(default_factory=list)
    extra_columns: list[ExtraColumn] = field(default_factory=list)

    def total_qty(self) -> int:
        total = 0
        for row in self.rows:
            try:
                total += int(str(row.qty).strip())
            except (TypeError, ValueError):
                continue
        return total


def _extra_value(row: InquiryRow, key: str) -> str:
    """动态补充列的单元格：有值就照常渲染，没值就标成「未提供」。"""
    cell = row.extra.get(key)
    if cell and (cell.en.strip() or cell.zh.strip()):
        return cell.render()
    return "Not provided\n未提供"


def _render_note(index: int, note: BiCell) -> str:
    """备注格式：Note N: <英文>\\n备注N：<中文>"""
    en, zh = (note.en or "").strip(), (note.zh or "").strip()
    if en and zh:
        return f"Note {index}: {en}\n备注{index}：{zh}"
    if en:
        return f"Note {index}: {en}"
    return f"备注{index}：{zh}"


def _title_text(table: InquiryTable) -> str:
    tail = ""
    if table.ref and table.country:
        tail = f" — Ref. {table.ref} ({table.country})"
    elif table.ref:
        tail = f" — Ref. {table.ref}"
    elif table.country:
        tail = f" — {table.country}"
    return f"VALVE INQUIRY / 阀门询价整理{tail}"


def export_inquiry(table: InquiryTable, out_path: str | Path) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()
    ws = wb.active
    ws.title = f"{table.ref} Inquiry" if table.ref else "Inquiry"

    # 标准 12 列 + AI 判定的「缺失必要参数」补充列
    all_columns = list(BASE_COLUMNS) + [
        (c.en, c.zh, c.width, c.align) for c in table.extra_columns
    ]
    extra_keys = [c.key for c in table.extra_columns]
    ncols = len(all_columns)

    # ---- 标题行 ----
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=ncols)
    title_cell = ws.cell(row=1, column=1, value=_title_text(table))
    title_cell.font = TITLE_FONT
    title_cell.fill = TITLE_FILL
    title_cell.alignment = WRAP_CENTER
    ws.row_dimensions[1].height = ROW_H_TITLE

    # ---- 表头行 ----
    for idx, (en, zh, _width, _align) in enumerate(all_columns, start=1):
        cell = ws.cell(row=2, column=idx, value=f"{en}\n{zh}" if en and zh else (en or zh))
        cell.font = HEAD_FONT
        cell.fill = HEAD_FILL
        cell.alignment = WRAP_CENTER
        cell.border = BORDER
    ws.row_dimensions[2].height = ROW_H_HEAD

    # ---- 数据行 ----
    r = 3
    for row in table.rows:
        values = row.to_cells() + [_extra_value(row, key) for key in extra_keys]
        for idx, value in enumerate(values, start=1):
            align = all_columns[idx - 1][3]
            cell = ws.cell(row=r, column=idx, value=value)
            cell.font = BODY_FONT
            cell.alignment = WRAP_CENTER if align == "center" else WRAP_LEFT
            cell.border = BORDER
        ws.row_dimensions[r].height = ROW_H_BODY
        r += 1

    # ---- 合计行 ----
    total = table.total_qty()
    if table.rows:
        ws.cell(row=r, column=1, value="Total\n合计")
        ws.cell(row=r, column=11, value=total)
        ws.cell(row=r, column=12, value=table.rows[-1].unit.render())
        for idx in range(1, ncols + 1):
            cell = ws.cell(row=r, column=idx)
            cell.font = TOTAL_FONT
            cell.fill = TOTAL_FILL
            cell.alignment = WRAP_CENTER
            cell.border = BORDER
        ws.row_dimensions[r].height = ROW_H_TOTAL
        r += 1

    # ---- 空行 ----
    r += 1

    # ---- 备注区 ----
    if table.notes:
        ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
        cell = ws.cell(row=r, column=1, value="REVIEW & CLARIFICATION NOTES / 审查与澄清备注")
        cell.font = NOTE_TITLE_FONT
        cell.fill = TITLE_FILL
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
        ws.row_dimensions[r].height = ROW_H_NOTE_TITLE
        r += 1

        for i, note in enumerate(table.notes, start=1):
            ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=ncols)
            cell = ws.cell(row=r, column=1, value=_render_note(i, note))
            cell.font = BODY_FONT
            cell.alignment = WRAP_LEFT
            ws.row_dimensions[r].height = ROW_H_NOTE
            r += 1

    # ---- 列宽 ----
    for idx, (_en, _zh, width, _align) in enumerate(all_columns, start=1):
        if width is None:
            continue
        ws.column_dimensions[ws.cell(row=2, column=idx).column_letter].width = width

    wb.save(out_path)
    return out_path
