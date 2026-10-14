"""本地脱敏：把真实项目的标识信息替换成代号。

**在把任何内容发往外部 AI 之前先跑这一步。**
映射表保存在本地，随时可以还原。

替换规则分两类：
  1. 整格替换 —— 位号 / P&ID 图号 / 管线号这类「本身就是标识」的列，整格换成 REF-NNNN
  2. 文本内替换 —— 描述、备注这类自由文本里的工程编号，按正则就地替换

用法：
    python src/masker.py <原始.xlsx> -o <脱敏后.xlsx> --map <映射表.json>
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook

# 工程编号：2位数字 + 3位字母 + 2位数字 + 2位字母 + 3位数字
# 例 11HAH01AA301 / 10HAH01CP001 / 11LAB11AA022 / 12LAB11AA021
REF_RE = re.compile(r"\b\d{2}[A-Z]{3}\d{2}[A-Z]{2}\d{3}\b")

# 抬头/文档编号这类固定串，直接列举替换
FIXED_RULES: dict[str, str] = {
    "LONG PHU 1 THERMAL POWER PLANT PROJECT": "[PROJECT NAME]",
    "LP108-00-PI-MTO-0002-01": "[DOC NO.]",
}

# 整格替换的列（表头区在 1-6 行，数据从第 7 行起）
WHOLE_CELL_COLUMNS = {
    2: "TAG",    # B  VALVE TAG NUMBER
    5: "PID",    # E  P&ID
    6: "LINE",   # F  PIPE LINE / EQP NO.
}

# 含自由文本、需要用正则就地替换的列
INLINE_TEXT_COLUMNS = {4, 24}   # D DESCRIPTION / X REMARKS

DATA_START_ROW = 7
HEADER_LAST_ROW = 6


class Masker:
    def __init__(self) -> None:
        self.mapping: dict[str, str] = {}
        self._seq = 0

    def code_for(self, raw: str) -> str:
        if raw in self.mapping:
            return self.mapping[raw]
        self._seq += 1
        code = f"REF-{self._seq:04d}"
        self.mapping[raw] = code
        return code

    def mask_text(self, text: str) -> str:
        out = text
        for src, dst in FIXED_RULES.items():
            out = out.replace(src, dst)
        return REF_RE.sub(lambda m: self.code_for(m.group(0)), out)

    def mask_cell_value(self, value):
        """整格替换：把这个值登记成一个代号。"""
        if value in (None, ""):
            return value
        text = str(value).strip()
        if not text:
            return value
        return self.code_for(text)


def mask_workbook(src: Path, dst: Path, map_path: Path) -> Masker:
    masker = Masker()
    wb = load_workbook(src)

    for ws in wb.worksheets:
        for row in ws.iter_rows():
            for cell in row:
                value = cell.value
                if value in (None, ""):
                    continue
                col = cell.column

                # 抬头区：固定规则 + 文本内编号
                if cell.row <= HEADER_LAST_ROW:
                    if isinstance(value, str):
                        new = masker.mask_text(value)
                        if new != value:
                            cell.value = new
                    continue

                # 数据区
                if col in WHOLE_CELL_COLUMNS:
                    cell.value = masker.mask_cell_value(value)
                elif col in INLINE_TEXT_COLUMNS and isinstance(value, str):
                    cell.value = masker.mask_text(value)

    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    wb.close()

    map_path.parent.mkdir(parents=True, exist_ok=True)
    payload = {
        "generated_at": datetime.now().isoformat(timespec="seconds"),
        "source": src.name,
        "masked": dst.name,
        "note": "原值 -> 代号。本文件只存本地，不要外发。",
        "mapping": masker.mapping,
    }
    map_path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8"
    )
    return masker


def restore_text(text: str, mapping: dict[str, str], table=("original", "masked")) -> str:
    """用映射表把代号还原回原值。"""
    reverse = {v: k for k, v in mapping.items()}
    out = text
    for code, original in sorted(reverse.items(), key=lambda kv: -len(kv[0])):
        out = out.replace(code, original)
    return out


def main() -> None:
    parser = argparse.ArgumentParser(description="本地脱敏")
    parser.add_argument("source", help="原始 xlsx")
    parser.add_argument("-o", "--out", required=True, help="脱敏后 xlsx 路径")
    parser.add_argument("--map", dest="map_path", required=True, help="映射表 json 路径")
    args = parser.parse_args()

    src = Path(args.source)
    if not src.exists():
        print(f"文件不存在：{src}")
        sys.exit(1)

    masker = mask_workbook(src, Path(args.out), Path(args.map_path))
    print(f"脱敏完成：{Path(args.out).name}")
    print(f"  替换条目：{len(masker.mapping)} 个")
    print(f"  映射表  ：{args.map_path}")
    print("  ⚠️ 映射表只存本地，不要外发。")


if __name__ == "__main__":
    main()
