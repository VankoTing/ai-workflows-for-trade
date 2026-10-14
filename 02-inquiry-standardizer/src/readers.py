"""多格式询价文件读取。

统一出口：

    doc = read_document(path)
    doc.kind      文件类型：excel / word / pdf / text / image
    doc.sheets    表格数据，每个 sheet 是 行 × 列 的二维数组
    doc.text      全文纯文本，用来喂给大模型审阅
    doc.warnings  读取过程中的提示（比如遇到图片需要 OCR）

设计原则：**能读出来的都读，读不出来的明确告知**，不要静默丢失内容。
"""

from __future__ import annotations

from dataclasses import dataclass, field
from pathlib import Path

EXCEL_EXT = {".xlsx", ".xlsm"}
WORD_EXT = {".docx"}
PDF_EXT = {".pdf"}
TEXT_EXT = {".txt", ".md", ".csv"}
IMAGE_EXT = {".png", ".jpg", ".jpeg", ".bmp", ".webp", ".tif", ".tiff"}


@dataclass
class Sheet:
    name: str
    rows: list[list[str]] = field(default_factory=list)

    def non_empty_rows(self) -> list[list[str]]:
        return [r for r in self.rows if any(c.strip() for c in r)]


@dataclass
class Document:
    path: Path
    kind: str
    sheets: list[Sheet] = field(default_factory=list)
    text: str = ""
    warnings: list[str] = field(default_factory=list)

    def to_text(self, max_chars: int = 20000) -> str:
        """给模型看的纯文本。表格按行展开，保留结构感。"""
        parts: list[str] = []
        for sheet in self.sheets:
            parts.append(f"### 表格：{sheet.name}")
            for i, row in enumerate(sheet.non_empty_rows(), 1):
                cells = [c for c in row]
                while cells and not cells[-1].strip():
                    cells.pop()
                if cells:
                    parts.append(f"{i}: " + " | ".join(cells))
        if self.text.strip():
            parts.append("### 正文")
            parts.append(self.text.strip())
        out = "\n".join(parts)
        if len(out) > max_chars:
            out = out[:max_chars] + f"\n…（已截断，原文共 {len(out)} 字符）"
        return out


def _norm_cell(value: object) -> str:
    if value is None:
        return ""
    return str(value).replace("\xa0", " ").strip()


# --------------------------------------------------------------------------
# Excel
# --------------------------------------------------------------------------

def read_excel(path: Path) -> Document:
    from openpyxl import load_workbook

    doc = Document(path=path, kind="excel")
    wb = load_workbook(path, data_only=True)
    for ws in wb.worksheets:
        sheet = Sheet(name=ws.title)
        for row in ws.iter_rows(values_only=True):
            sheet.rows.append([_norm_cell(c) for c in row])
        # 去掉完全空白的尾部行
        while sheet.rows and not any(c.strip() for c in sheet.rows[-1]):
            sheet.rows.pop()
        doc.sheets.append(sheet)

    if not any(s.non_empty_rows() for s in doc.sheets):
        doc.warnings.append("表格里没有读到内容，可能是纯图片型表格，需要 OCR。")
    wb.close()
    return doc


# --------------------------------------------------------------------------
# Word
# --------------------------------------------------------------------------

def read_word(path: Path) -> Document:
    import docx

    doc = Document(path=path, kind="word")
    d = docx.Document(str(path))

    for idx, table in enumerate(d.tables, 1):
        sheet = Sheet(name=f"Table {idx}")
        for row in table.rows:
            sheet.rows.append([_norm_cell(c.text) for c in row.cells])
        doc.sheets.append(sheet)

    paragraphs = [_norm_cell(p.text) for p in d.paragraphs]
    doc.text = "\n".join(p for p in paragraphs if p)

    if not doc.sheets and not doc.text:
        doc.warnings.append("Word 文档里没读到内容。")
    return doc


# --------------------------------------------------------------------------
# PDF
# --------------------------------------------------------------------------

def read_pdf(path: Path) -> Document:
    from pypdf import PdfReader

    doc = Document(path=path, kind="pdf")
    reader = PdfReader(str(path))
    pages: list[str] = []
    for i, page in enumerate(reader.pages, 1):
        try:
            content = page.extract_text() or ""
        except Exception as exc:  # noqa: BLE001
            doc.warnings.append(f"第 {i} 页解析失败：{exc}")
            content = ""
        if content.strip():
            pages.append(f"--- 第 {i} 页 ---\n{content.strip()}")
    doc.text = "\n\n".join(pages)

    if not doc.text.strip():
        doc.warnings.append(
            "PDF 里没有可提取的文字，很可能是扫描件，需要 OCR 或改用图片通道。"
        )
    return doc


# --------------------------------------------------------------------------
# 纯文本
# --------------------------------------------------------------------------

def read_text(path: Path) -> Document:
    doc = Document(path=path, kind="text")
    for enc in ("utf-8-sig", "utf-8", "gbk", "latin-1"):
        try:
            doc.text = path.read_text(encoding=enc)
            break
        except UnicodeDecodeError:
            continue
    else:
        doc.warnings.append("文本编码无法识别。")
    return doc


# --------------------------------------------------------------------------
# 图片（占位）
# --------------------------------------------------------------------------

def read_image(path: Path) -> Document:
    doc = Document(path=path, kind="image")
    doc.warnings.append(
        "这是一张图片，需要 OCR 才能取出文字。第一版先不处理，"
        "建议先让客户提供 Excel / Word / 文字型 PDF。"
    )
    return doc


# --------------------------------------------------------------------------
# 入口
# --------------------------------------------------------------------------

def read_document(path: str | Path) -> Document:
    p = Path(path)
    if not p.exists():
        raise FileNotFoundError(f"文件不存在：{p}")

    ext = p.suffix.lower()
    if ext in EXCEL_EXT:
        return read_excel(p)
    if ext in WORD_EXT:
        return read_word(p)
    if ext in PDF_EXT:
        return read_pdf(p)
    if ext in IMAGE_EXT:
        return read_image(p)
    if ext in TEXT_EXT:
        return read_text(p)

    # 兜底：尝试按 Excel 打开（有些导出文件后缀不规范）
    try:
        return read_excel(p)
    except Exception:  # noqa: BLE001
        doc = read_text(p)
        doc.warnings.append(f"未识别的文件类型 {ext}，已按纯文本尝试读取。")
        return doc
